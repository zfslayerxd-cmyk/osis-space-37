import os, json, io, csv, uuid, secrets, hmac, hashlib, datetime, logging, zipfile
from pathlib import Path
from functools import wraps
import psycopg
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from flask import Flask, request, session, jsonify, send_file, send_from_directory, g
from werkzeug.exceptions import HTTPException
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from core import DIVS, ROLES, hashpw, checkpw, admin, can_read, can_write, clean
ROOT=Path(__file__).resolve().parent
KINDS=json.loads((ROOT/'kinds.json').read_text(encoding='utf-8'))
XL={'finance':{'tanggal':'date','jenis':'type','nominal':'amount','kategori':'category','keterangan':'description'},'tasks':{'nama_program':'title','divisi':'division','penanggung_jawab':'owner','target_selesai':'deadline','status':'status','deskripsi':'description'}}
app=Flask(__name__,static_folder=str(ROOT/'public/static'),static_url_path='/static')
app.config.update(SECRET_KEY=os.environ.get('SECRET_KEY'),MAX_CONTENT_LENGTH=3*1024*1024,SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Lax',SESSION_COOKIE_SECURE=bool(os.environ.get('VERCEL')),PERMANENT_SESSION_LIFETIME=datetime.timedelta(hours=8))
def db():
    if 'db' not in g:g.db=psycopg.connect(os.environ['DATABASE_URL'],row_factory=dict_row,connect_timeout=8)
    return g.db
@app.teardown_appcontext
def close_db(error):
    c=g.pop('db',None)
    if c:c.close()
def query(sql,args=(),many=False):
    cur=db().execute(sql,args)
    return cur.fetchall() if many else cur.fetchone()
def write(sql,args=()):
    c=db();cur=c.execute(sql,args);c.commit();return cur

def public(u,full=False):
    result={k:u[k] for k in ('id','name','kelas','jabatan','role','division')}
    if full:result.update(username=u['username'],must_change=u['must_change'])
    return result

def auth(fn):
    @wraps(fn)
    def wrapped(*args,**kwargs):
        u=query('SELECT * FROM osis_users WHERE id=%s',(session.get('uid',-1),))
        if not u or u['pwv']!=session.get('pwv'):return jsonify(error='Sesi berakhir. Login kembali.',code='auth'),401
        if u['must_change'] and fn.__name__!='password':return jsonify(error='Ganti password sementara dahulu.',code='change'),403
        return fn(u,*args,**kwargs)
    return wrapped

def body():
    value=request.get_json(silent=True)
    if not isinstance(value,dict):raise ValueError('Kirim data JSON berupa objek.')
    return value

def div_for(k,v):return KINDS[k].get('division') or v.get('division','')
def mod(u,r):return can_write(u,r['kind'],r['division'],KINDS) and (u['role']!='anggota' or r['creator']==u['id'])
def docread(u):return u['role'] in ('pembina','ketua','wakil','sekretaris')
def docwrite(u):return u['role'] in ('pembina','ketua','sekretaris')
@app.before_request
def protect():
    if request.path.startswith('/api/') and request.path!='/api/health':
        if not os.environ.get('SECRET_KEY') or not os.environ.get('DATABASE_URL'):return jsonify(error='DATABASE_URL atau SECRET_KEY belum tersedia. Atur environment lalu redeploy.'),503
        if request.method in ('POST','PUT','DELETE','PATCH'):
            token=request.headers.get('X-CSRF-Token','');expected=session.get('csrf','')
            if not expected or not hmac.compare_digest(token,expected):return jsonify(error='Token sesi tidak valid. Muat ulang halaman.'),403
@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Referrer-Policy']='same-origin'
    response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"
    if request.path.startswith('/api/'):response.headers['Cache-Control']='no-store'
    return response
@app.errorhandler(Exception)
def errors(error):
    if isinstance(error,HTTPException):return jsonify(error=error.description),error.code
    if isinstance(error,(ValueError,TypeError)):return jsonify(error=str(error)),400
    if isinstance(error,psycopg.errors.UniqueViolation):return jsonify(error='Data duplikat. Gunakan nilai lain.'),409
    logging.exception('OSIS server error')
    return jsonify(error='Server gagal memproses permintaan. Periksa Runtime Logs dan koneksi/schema database.'),500
@app.get('/')
def index():return send_from_directory(ROOT/'templates','index.html')
@app.get('/api/health')
def health():return jsonify(ok=True,service='osis-space',version='2.0.0')
@app.get('/api/session')
def current():
    session.setdefault('csrf',secrets.token_hex(24))
    u=query('SELECT * FROM osis_users WHERE id=%s',(session.get('uid',-1),)) if session.get('uid') else None
    if u and u['pwv']!=session.get('pwv'):u=None
    return jsonify(csrf=session['csrf'],user=public(u,True) if u else None,needs_setup=not bool(query('SELECT id FROM osis_users LIMIT 1')))
@app.post('/api/setup')
def setup():
    expected=os.environ.get('SETUP_TOKEN','');got=body().get('token')
    if not expected or not isinstance(got,str) or not hmac.compare_digest(got,expected):return jsonify(error='Token setup salah.'),403
    seed=json.loads((ROOT/'seed_accounts.json').read_text(encoding='utf-8'));c=db()
    c.execute('SELECT pg_advisory_xact_lock(8651031)')
    if c.execute('SELECT id FROM osis_users LIMIT 1').fetchone():c.rollback();return jsonify(error='Database sudah berisi akun. Setup tidak menimpa akun lama.'),409
    for s in seed:c.execute('INSERT INTO osis_users(username,password,name,kelas,jabatan,role,division) VALUES(%s,%s,%s,%s,%s,%s,%s)',tuple(s[k] for k in ('username','password','name','kelas','jabatan','role','division')))
    c.commit();return jsonify(ok=True,count=len(seed))
@app.post('/api/login')
def login():
    d=body();un=str(d.get('username','')).strip().lower();pw=str(d.get('password',''))
    if not un or len(un)>80 or len(pw)>200:return jsonify(error='Input login tidak valid.'),400
    key=hashlib.sha256(((request.remote_addr or 'unknown')+'|'+un).encode()).hexdigest()
    a=query('SELECT * FROM osis_attempts WHERE k=%s',(key,))
    if a and a['count']>=8 and a['started']>datetime.datetime.now(datetime.timezone.utc)-datetime.timedelta(minutes=15):return jsonify(error='Terlalu banyak percobaan. Tunggu 15 menit.'),429
    u=query('SELECT * FROM osis_users WHERE username=%s',(un,))
    if not u or not checkpw(u['password'],pw):
        write("INSERT INTO osis_attempts(k,count,started) VALUES(%s,1,now()) ON CONFLICT(k) DO UPDATE SET count=CASE WHEN osis_attempts.started<now()-interval '15 minutes' THEN 1 ELSE osis_attempts.count+1 END,started=CASE WHEN osis_attempts.started<now()-interval '15 minutes' THEN now() ELSE osis_attempts.started END",(key,))
        return jsonify(error='Username atau password salah.'),401
    write('DELETE FROM osis_attempts WHERE k=%s',(key,));session.clear();session.update(uid=u['id'],pwv=u['pwv'],csrf=secrets.token_hex(24));session.permanent=True
    return jsonify(ok=True,csrf=session['csrf'])
@app.post('/api/logout')
def logout():session.clear();return jsonify(ok=True)
@app.post('/api/password')
@auth
def password(u):
    d=body();new=str(d.get('new',''))
    if not checkpw(u['password'],str(d.get('current',''))):return jsonify(error='Password lama salah.'),400
    if not 8<=len(new)<=200 or checkpw(u['password'],new):return jsonify(error='Gunakan password baru berbeda, minimal 8 karakter.'),400
    write('UPDATE osis_users SET password=%s,must_change=false,pwv=pwv+1 WHERE id=%s',(hashpw(new),u['id']));session['pwv']=u['pwv']+1
    return jsonify(ok=True)
@app.get('/api/meta')
@auth
def meta(u):
    s=query('SELECT value FROM osis_settings WHERE id=1')
    if not s:raise RuntimeError('osis_settings id=1 missing')
    return jsonify(user=public(u,True),admin=admin(u),kinds=KINDS,divs=DIVS,roles=ROLES,perms={k:{'read':can_read(u,k,KINDS),'write':can_write(u,k,u['division'],KINDS)} for k in KINDS},docs={'read':docread(u),'write':docwrite(u)},settings=s['value'])
@app.route('/api/users',methods=['GET','POST'])
@auth
def users(u):
    if request.method=='GET':return jsonify(items=[public(r,admin(u)) for r in query('SELECT * FROM osis_users ORDER BY id',many=True)])
    if not admin(u):return jsonify(error='Tidak berwenang.'),403
    d=body();un=str(d.get('username','')).strip().lower();name=str(d.get('name','')).strip()
    if not un or len(un)>40 or not all(c.isascii() and (c.isalnum() or c in '._-') for c in un) or not name or len(name)>80 or d.get('role') not in ROLES or d.get('division') not in DIVS+['Pembina']:raise ValueError('Data akun tidak valid.')
    pw=secrets.token_urlsafe(12)
    write('INSERT INTO osis_users(username,password,name,kelas,jabatan,role,division) VALUES(%s,%s,%s,%s,%s,%s,%s)',(un,hashpw(pw),name,str(d.get('kelas',''))[:10],str(d.get('jabatan',''))[:80],d['role'],d['division']))
    return jsonify(username=un,password=pw)
@app.post('/api/users/<int:uid>/reset')
@auth
def reset(u,uid):
    if not admin(u) or uid==u['id']:return jsonify(error='Tidak berwenang. Ganti password sendiri lewat Akun saya.'),403
    target=query('SELECT * FROM osis_users WHERE id=%s',(uid,))
    if not target:return jsonify(error='Akun tidak ditemukan.'),404
    pw=secrets.token_urlsafe(12);write('UPDATE osis_users SET password=%s,must_change=true,pwv=pwv+1 WHERE id=%s',(hashpw(pw),uid))
    return jsonify(username=target['username'],password=pw)
@app.put('/api/settings')
@auth
def settings(u):
    if not admin(u):return jsonify(error='Tidak berwenang.'),403
    d=body();v={k:str(d.get(k,'')).strip() for k in ('school','period','tagline')}
    if any(not x or len(x)>150 for x in v.values()):raise ValueError('Isi semua kolom, maksimal 150 karakter.')
    write('UPDATE osis_settings SET value=%s WHERE id=1',(Jsonb(v),));return jsonify(ok=True)
@app.route('/api/records/<k>',methods=['GET','POST'])
@auth
def records(u,k):
    if not can_read(u,k,KINDS):return jsonify(error='Tidak punya akses.'),403
    if request.method=='GET':
        items=[]
        for r in query('SELECT * FROM osis_records WHERE kind=%s ORDER BY created DESC LIMIT 2001',(k,),many=True):items.append({**r['value'],'id':r['id'],'_can':mod(u,r)})
        return jsonify(items=items[:2000],truncated=len(items)>2000)
    v=clean(k,body(),KINDS);div=div_for(k,v)
    if not can_write(u,k,div,KINDS):return jsonify(error='Tidak boleh menambah data divisi ini.'),403
    write('INSERT INTO osis_records(id,kind,division,value,creator) VALUES(%s,%s,%s,%s,%s)',(uuid.uuid4().hex,k,div,Jsonb(v),u['id']))
    return jsonify(ok=True)
@app.route('/api/records/<k>/<rid>',methods=['PUT','DELETE'])
@auth
def record(u,k,rid):
    if k not in KINDS:return jsonify(error='Modul tidak ditemukan.'),404
    r=query('SELECT * FROM osis_records WHERE id=%s AND kind=%s',(rid,k))
    if not r or not mod(u,r):return jsonify(error='Tidak berwenang.'),403
    if request.method=='DELETE':write('DELETE FROM osis_records WHERE id=%s',(rid,));return jsonify(ok=True)
    v=clean(k,body(),KINDS);div=div_for(k,v)
    if not can_write(u,k,div,KINDS):return jsonify(error='Tidak berwenang ke divisi tujuan.'),403
    write('UPDATE osis_records SET division=%s,value=%s WHERE id=%s',(div,Jsonb(v),rid));return jsonify(ok=True)
@app.get('/api/export/finance')
@auth
def export(u):
    text=io.StringIO();writer=csv.writer(text);writer.writerow(XL['finance'])
    for r in query("SELECT value FROM osis_records WHERE kind='finance' ORDER BY created DESC",many=True):writer.writerow([("'"+str(v)) if str(v).startswith(('=','+','-','@','\t','\r')) else v for v in (r['value'].get(n,'') for n in XL['finance'].values())])
    return send_file(io.BytesIO(('\ufeff'+text.getvalue()).encode('utf-8')),mimetype='text/csv',as_attachment=True,download_name='keuangan-osis.csv')
@app.route('/api/documents',methods=['GET','POST'])
@auth
def documents(u):
    if not docread(u):return jsonify(error='Tidak berwenang.'),403
    if request.method=='GET':return jsonify(items=query('SELECT id,name,size,created::text FROM osis_documents ORDER BY created DESC',many=True))
    if not docwrite(u):return jsonify(error='Tidak berwenang.'),403
    f=request.files.get('file')
    if not f or not (f.filename or '').lower().endswith('.pdf'):raise ValueError('Pilih PDF.')
    content=f.read()
    if not content.startswith(b'%PDF-'):raise ValueError('Berkas PDF tidak valid.')
    write('INSERT INTO osis_documents(id,name,size,content) VALUES(%s,%s,%s,%s)',(uuid.uuid4().hex,(secure_filename(f.filename) or 'dokumen.pdf')[:250],len(content),content))
    return jsonify(ok=True)
@app.route('/api/documents/<rid>',methods=['GET','DELETE'])
@auth
def document(u,rid):
    if not docread(u):return jsonify(error='Tidak berwenang.'),403
    if request.method=='DELETE':
        if not docwrite(u):return jsonify(error='Tidak berwenang.'),403
        write('DELETE FROM osis_documents WHERE id=%s',(rid,));return jsonify(ok=True)
    r=query('SELECT * FROM osis_documents WHERE id=%s',(rid,))
    if not r:return jsonify(error='PDF tidak ditemukan.'),404
    return send_file(io.BytesIO(bytes(r['content'])),as_attachment=True,download_name=r['name'],mimetype='application/pdf')
@app.get('/api/template/<k>')
@auth
def template(u,k):
    if k not in XL or not any(can_write(u,k,d,KINDS) for d in DIVS):return jsonify(error='Tidak berwenang.'),403
    w=Workbook();w.active.append(list(XL[k]));b=io.BytesIO();w.save(b);b.seek(0)
    return send_file(b,as_attachment=True,download_name=k+'.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
@app.post('/api/excel')
@auth
def excel(u):
    k=request.form.get('kind');f=request.files.get('file');commit=request.form.get('commit')=='1'
    if k not in XL or not any(can_write(u,k,d,KINDS) for d in DIVS):return jsonify(error='Tidak berwenang.'),403
    if not f or not (f.filename or '').lower().endswith('.xlsx'):raise ValueError('Pilih Excel .xlsx.')
    content=f.read()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if len(archive.infolist())>2000 or sum(x.file_size for x in archive.infolist())>30*1024*1024:raise ValueError('Isi Excel terlalu besar.')
    w=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
    try:
        rows=w.active.iter_rows(values_only=True);heads=[str(v or '').strip().lower() for v in next(rows,[])]
        if len(heads)!=len(set(heads)) or not set(XL[k]).issubset(heads):raise ValueError('Header harus mengikuti template.')
        valid=[];errors=[]
        for i,row in enumerate(rows,2):
            if i>2001:raise ValueError('Maksimal 2000 baris.')
            if all(v is None for v in row):continue
            try:
                d={}
                for h,n in XL[k].items():
                    idx=heads.index(h);v=row[idx] if idx<len(row) else None
                    if isinstance(v,datetime.datetime):v=v.date().isoformat()
                    elif isinstance(v,datetime.date):v=v.isoformat()
                    if isinstance(v,str) and v.startswith('='):raise ValueError('Rumus tidak didukung.')
                    d[n]=v
                v=clean(k,d,KINDS)
                if not can_write(u,k,div_for(k,v),KINDS):raise ValueError('Divisi tidak berwenang.')
                valid.append(v)
            except (ValueError,TypeError) as e:errors.append({'row':i,'message':str(e)})
    finally:w.close()
    added=0
    if commit:
        if errors or not valid:raise ValueError('Perbaiki seluruh kesalahan sebelum impor.')
        c=db();c.execute('SELECT pg_advisory_xact_lock(8651032)')
        known={json.dumps(r['value'],sort_keys=True) for r in c.execute('SELECT value FROM osis_records WHERE kind=%s',(k,)).fetchall()}
        for v in valid:
            sig=json.dumps(v,sort_keys=True)
            if sig in known:continue
            known.add(sig);c.execute('INSERT INTO osis_records(id,kind,division,value,creator) VALUES(%s,%s,%s,%s,%s)',(uuid.uuid4().hex,k,div_for(k,v),Jsonb(v),u['id']));added+=1
        c.commit()
    return jsonify(rows=valid[:30],total=len(valid),errors=errors[:50],error_count=len(errors),added=added,skipped=len(valid)-added if commit else 0)
if __name__=='__main__':app.run(host='127.0.0.1',port=int(os.environ.get('PORT','8000')),debug=False)
